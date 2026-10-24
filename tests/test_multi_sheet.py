"""A workbook with a Sales sheet and an Expenses sheet is read in full: the first sheet with the connector's mapping, the
others with the same column rules the Data page suggests. A sheet that cannot be understood is reported, never fatal."""
import datetime as dt
import io

import openpyxl

from app.connectors.file_connector import extra_sheet_events, other_sheet_names, parse_file_to_dataframe
from tests.test_block_c import make_owner, mk


def _workbook(day, extra=True, junk=False):
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Sales"
    ws.append(["Date", "Receipt", "Branch", "Amount"])
    ws.append([day, "R1", "Meru", 1000]); ws.append([day, "R2", "Meru", 500])
    if extra:
        ex = wb.create_sheet("Expenses")
        ex.append(["Date", "Reference", "Category", "Expense Amount"])
        ex.append([day, "E1", "Rent", 300]); ex.append([day, "E2", "Salaries", 200])
    if junk:
        wb.create_sheet("Notes").append(["Remember", "to", "order"])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def test_first_sheet_only_workbooks_behave_exactly_as_before():
    content = _workbook(dt.datetime(2026, 10, 1), extra=False)
    assert other_sheet_names(content, ".xlsx") == [] and extra_sheet_events(content, ".xlsx", "t", "c", "s") == ([], [], [])
    assert list(parse_file_to_dataframe(content, ".xlsx").columns) == ["Date", "Receipt", "Branch", "Amount"]


def test_expenses_sheet_is_read_as_expenses():
    events, skipped, notes = extra_sheet_events(_workbook(dt.datetime(2026, 10, 1)), ".xlsx", "t", "c", "s")
    assert {e["entity"] for e in events} == {"expense"} and sorted(e["amount"] for e in events) == [200.0, 300.0] and not skipped
    assert notes == [{"sheet": "Expenses", "status": "read", "as": "expense", "rows": 2, "skipped": 0}]


def test_unreadable_extra_sheet_is_reported_and_does_not_stop_the_rest():
    events, _, notes = extra_sheet_events(_workbook(dt.datetime(2026, 10, 1), junk=True), ".xlsx", "t", "c", "s")
    assert len(events) == 2 and [n["status"] for n in notes] == ["read", "left out"]


def test_upload_counts_sales_and_expenses_from_both_sheets(client):
    t, h = make_owner(client, "multisheet")
    cid = mk(client, t, h, connector_type="file", config={"entity": "sale", "event_type": "sale.created", "timestamp_col": "Date",
             "amount_col": "Amount", "entity_id_col": "Receipt"}).json()["id"]
    day = dt.datetime.utcnow() - dt.timedelta(days=2)
    r = client.post(f"/api/connectors/{cid}/upload", data={"tenant_id": t}, headers=h,
                    files={"file": ("book.xlsx", _workbook(day), "application/octet-stream")})
    assert r.status_code == 200 and r.json()["rows_processed"] == 4 and r.json()["other_sheets"][0]["status"] == "read", r.text
    k = client.get(f"/api/questions/kpis?tenant_id={t}&period_days=30", headers=h).json()
    assert k["revenue"]["current"] == 1500 and k["expenses"]["current"] == 500
    again = client.post(f"/api/connectors/{cid}/upload", data={"tenant_id": t}, headers=h,
                        files={"file": ("book.xlsx", _workbook(day), "application/octet-stream")})
    k2 = client.get(f"/api/questions/kpis?tenant_id={t}&period_days=30", headers=h).json()
    assert again.status_code == 200 and k2["expenses"]["current"] == 500       # uploading the same book twice never double counts
